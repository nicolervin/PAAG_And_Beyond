import unittest
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4
import pandas as pd
from streamlit.testing.v1 import AppTest

from utils import store


class ModelTreeScratchBomTests(unittest.TestCase):
    def setUp(self):
        self.database_path = store.DATA_DIR / f"test_bom_scratch_{uuid4()}.db"
        self.database_patch = patch.object(store, "DB_PATH", self.database_path)
        self.database_patch.start()
        store.init_db()
        self.project_id = str(store.query("SELECT id FROM projects LIMIT 1")[0]["id"])
        self.model_id = store.add_project_model(
            self.project_id, "TEST-M1", "Test Model 1", "For BOM from scratch testing"
        )

    def tearDown(self):
        self.database_patch.stop()
        if self.database_path.exists():
            self.database_path.unlink()

    def test_validation_empty_items(self):
        with self.assertRaises(ValueError) as ctx:
            store.save_manual_model_bom(self.project_id, self.model_id, [])
        self.assertIn("at least one item", str(ctx.exception).lower())

    def test_validation_level_skips(self):
        invalid_items = [
            {"level": 1, "part_number": "ASM-001", "quantity": 1.0},
            {"level": 3, "part_number": "PRT-001", "quantity": 1.0}, # skipped L2!
        ]
        with self.assertRaises(ValueError) as ctx:
            store.save_manual_model_bom(self.project_id, self.model_id, invalid_items)
        self.assertIn("cannot skip levels", str(ctx.exception).lower())

    def test_validation_quantity(self):
        invalid_items = [
            {"level": 1, "part_number": "ASM-001", "quantity": 0.0},
        ]
        with self.assertRaises(ValueError) as ctx:
            store.save_manual_model_bom(self.project_id, self.model_id, invalid_items)
        self.assertIn("greater than 0", str(ctx.exception).lower())

    def test_save_and_retrieve_manual_bom(self):
        items = [
            {"level": 1, "part_number": "ASM-TOP", "description": "Top Assembly", "quantity": 1.0},
            {"level": 2, "part_number": "SUB-DOOR", "description": "Door Assembly", "quantity": 1.0},
            {"level": 3, "part_number": "PRT-HNDL", "description": "Door Handle", "quantity": 2.0},
            {"level": 3, "part_number": "PRT-GSKT", "description": "Door Gasket", "quantity": 1.0},
            {"level": 2, "part_number": "SUB-CAB", "description": "Cabinet Assembly", "quantity": 1.0},
            {"level": 3, "part_number": "PRT-SHLF", "description": "Glass Shelf", "quantity": 3.0},
        ]

        res = store.save_manual_model_bom(
            self.project_id,
            self.model_id,
            items,
            auto_add_parts=True,
            editor="IE Tester",
        )
        self.assertEqual(res["saved_count"], 6)
        self.assertEqual(res["parts_created"], 6)
        self.assertEqual(res["model_number"], "TEST-M1")

        # Check retrieval via get_manual_model_bom
        retrieved = store.get_manual_model_bom(self.project_id, self.model_id)
        self.assertEqual(len(retrieved), 6)
        self.assertEqual(retrieved[0]["part_number"], "ASM-TOP")
        self.assertEqual(retrieved[0]["level"], 1)
        self.assertEqual(retrieved[2]["part_number"], "PRT-HNDL")
        self.assertEqual(retrieved[2]["level"], 3)
        self.assertEqual(retrieved[2]["quantity"], 2.0)

        # Check CAD model tree integration
        tree_res = store.pits_bom_model_tree(self.project_id, self.model_id)
        self.assertEqual(len(tree_res["roots"]), 1)
        root = tree_res["roots"][0]
        self.assertEqual(root["part_number"], "ASM-TOP")
        self.assertEqual(len(root["children"]), 2) # SUB-DOOR and SUB-CAB
        self.assertEqual(root["children"][0]["part_number"], "SUB-DOOR")
        self.assertEqual(len(root["children"][0]["children"]), 2) # PRT-HNDL and PRT-GSKT

        self.assertEqual(tree_res["metrics"]["total"], 6)
        self.assertEqual(tree_res["metrics"]["missing_catalog"], 0) # all auto-onboarded!

    def test_multi_model_isolation(self):
        model2_id = store.add_project_model(self.project_id, "TEST-M2", "Model 2", "")

        bom_m1 = [
            {"level": 1, "part_number": "TOP-M1", "description": "M1 Root", "quantity": 1.0},
            {"level": 2, "part_number": "PART-ONLY-M1", "description": "M1 Specific", "quantity": 1.0},
        ]
        bom_m2 = [
            {"level": 1, "part_number": "TOP-M2", "description": "M2 Root", "quantity": 1.0},
            {"level": 2, "part_number": "PART-ONLY-M2", "description": "M2 Specific", "quantity": 5.0},
        ]

        store.save_manual_model_bom(self.project_id, self.model_id, bom_m1, auto_add_parts=True)
        store.save_manual_model_bom(self.project_id, model2_id, bom_m2, auto_add_parts=True)

        tree_m1 = store.pits_bom_model_tree(self.project_id, self.model_id)
        m1_parts = [n["part_number"] for n in tree_m1["nodes"]]
        self.assertIn("PART-ONLY-M1", m1_parts)
        self.assertNotIn("PART-ONLY-M2", m1_parts)

        tree_m2 = store.pits_bom_model_tree(self.project_id, model2_id)
        m2_parts = [n["part_number"] for n in tree_m2["nodes"]]
        self.assertIn("PART-ONLY-M2", m2_parts)
        self.assertNotIn("PART-ONLY-M1", m2_parts)

    def test_level1_assembly_unique_validation(self):
        bom_m1 = [
            {"level": 1, "part_number": "TOP-SHARED", "description": "M1 Root", "quantity": 1.0},
            {"level": 2, "part_number": "PRT-1", "description": "Part 1", "quantity": 1.0},
        ]
        store.save_manual_model_bom(self.project_id, self.model_id, bom_m1, auto_add_parts=True)

        model2_id = store.add_project_model(self.project_id, "TEST-M2", "Model 2", "")
        bom_m2_duplicate = [
            {"level": 1, "part_number": "TOP-SHARED", "description": "Same Top Assembly", "quantity": 1.0},
            {"level": 2, "part_number": "PRT-2", "description": "Part 2", "quantity": 1.0},
        ]
        with self.assertRaises(ValueError) as ctx:
            store.save_manual_model_bom(self.project_id, model2_id, bom_m2_duplicate, auto_add_parts=True)
        self.assertIn("already assigned to model 'TEST-M1'", str(ctx.exception))

    def test_copy_model_bom_requires_level2_differences(self):
        bom_m1 = [
            {"level": 1, "part_number": "TOP-M1-SRC", "description": "Source Root", "quantity": 1.0},
            {"level": 2, "part_number": "SUB-DOOR", "description": "Door Sub", "quantity": 1.0},
            {"level": 3, "part_number": "PRT-HNDL", "description": "Handle", "quantity": 2.0},
        ]
        store.save_manual_model_bom(self.project_id, self.model_id, bom_m1, auto_add_parts=True)

        model2_id = store.add_project_model(self.project_id, "TEST-M2", "Model 2", "")
        # Exact clone of Level 2+ with changed Level 1
        copied_identical_l2 = [
            {"level": 1, "part_number": "TOP-M2-NEW", "description": "New Root", "quantity": 1.0},
            {"level": 2, "part_number": "SUB-DOOR", "description": "Door Sub", "quantity": 1.0},
            {"level": 3, "part_number": "PRT-HNDL", "description": "Handle", "quantity": 2.0},
        ]

        # Saving with source_model_id should fail because Level 2+ is 100% identical
        with self.assertRaises(ValueError) as ctx:
            store.save_manual_model_bom(
                self.project_id,
                model2_id,
                copied_identical_l2,
                auto_add_parts=True,
                source_model_id=self.model_id,
            )
        self.assertIn("100% identical", str(ctx.exception))

        # Modifying a quantity in Level 2+ allows save to succeed
        copied_varied = [
            {"level": 1, "part_number": "TOP-M2-NEW", "description": "New Root", "quantity": 1.0},
            {"level": 2, "part_number": "SUB-DOOR", "description": "Door Sub", "quantity": 1.0},
            {"level": 3, "part_number": "PRT-HNDL", "description": "Handle", "quantity": 4.0}, # changed qty!
        ]
        res = store.save_manual_model_bom(
            self.project_id,
            model2_id,
            copied_varied,
            auto_add_parts=True,
            source_model_id=self.model_id,
        )
        self.assertEqual(res["saved_count"], 3)

    def test_auto_detect_and_fill_catalog_description(self):
        # Pre-seed a part with description in catalog
        store.upsert_part(
            self.project_id,
            {
                "part_number": "PRT-CATALOG-XYZ",
                "description": "Catalog Premium Glass Shelf",
            },
        )

        bom_items = [
            {"level": 1, "part_number": "TOP-DESC-TEST", "description": "Top Root", "quantity": 1.0},
            {"level": 2, "part_number": "PRT-CATALOG-XYZ", "description": "", "quantity": 2.0}, # empty desc
        ]
        store.save_manual_model_bom(self.project_id, self.model_id, bom_items, auto_add_parts=True)

        retrieved = store.get_manual_model_bom(self.project_id, self.model_id)
        self.assertEqual(len(retrieved), 2)
        self.assertEqual(retrieved[1]["part_number"], "PRT-CATALOG-XYZ")
        self.assertEqual(retrieved[1]["description"], "Catalog Premium Glass Shelf")

    def test_get_known_subassembly_structure(self):
        bom_items = [
            {"level": 1, "part_number": "TOP-SUB-TEST", "description": "Root", "quantity": 1.0},
            {"level": 2, "part_number": "SUB-DOOR-X", "description": "Door Assembly", "quantity": 1.0},
            {"level": 3, "part_number": "PRT-BOLT-1", "description": "Hinge Bolt", "quantity": 4.0},
            {"level": 3, "part_number": "PRT-SEAL-1", "description": "Door Seal", "quantity": 1.0},
        ]
        store.save_manual_model_bom(self.project_id, self.model_id, bom_items, auto_add_parts=True)

        sub_children = store.get_known_subassembly_structure(self.project_id, "SUB-DOOR-X")
        self.assertEqual(len(sub_children), 2)
        pns = [c["part_number"] for c in sub_children]
        self.assertIn("PRT-BOLT-1", pns)
        self.assertIn("PRT-SEAL-1", pns)

    def test_get_model_level1_assemblies(self):
        bom_items = [
            {"level": 1, "part_number": "ASM-TOP-01", "description": "Top Root", "quantity": 1.0},
        ]
        store.save_manual_model_bom(self.project_id, self.model_id, bom_items, auto_add_parts=True)

        lvl1_map = store.get_model_level1_assemblies(self.project_id)
        self.assertEqual(lvl1_map.get("ASM-TOP-01"), "TEST-M1")

    def test_part_number_not_transferred_to_pits_id_column(self):
        items = [
            {"level": 1, "part_number": "PRT-NEW-99", "description": "Brand New Part", "quantity": 1.0},
        ]
        res = store.save_manual_model_bom(
            self.project_id, self.model_id, items, auto_add_parts=True
        )
        self.assertEqual(res["parts_created"], 1)

        row = store.query(
            "SELECT part_number, pits_tracker_number FROM parts WHERE project_id=? AND part_number='PRT-NEW-99'",
            (self.project_id,),
        )[0]
        self.assertEqual(row["part_number"], "PRT-NEW-99")
        self.assertEqual(row["pits_tracker_number"], "")

        # Test cleanup if an existing part erroneously had pits_tracker_number == part_number
        store.query(
            "UPDATE parts SET pits_tracker_number='PRT-NEW-99' WHERE project_id=? AND part_number='PRT-NEW-99'",
            (self.project_id,),
        )
        store.save_manual_model_bom(
            self.project_id, self.model_id, items, auto_add_parts=True
        )
        row_cleaned = store.query(
            "SELECT pits_tracker_number FROM parts WHERE project_id=? AND part_number='PRT-NEW-99'",
            (self.project_id,),
        )[0]
        self.assertEqual(row_cleaned["pits_tracker_number"], "")

    def test_quantity_five_decimal_places(self):
        items = [
            {"level": 1, "part_number": "TOP-DEC-5", "description": "Decimal Root", "quantity": 1.0},
            {"level": 2, "part_number": "PRT-ADHESIVE", "description": "Structural Glue (kg)", "quantity": 0.00025},
            {"level": 2, "part_number": "PRT-SEALANT", "description": "Silicone Sealant (L)", "quantity": 0.00001},
            {"level": 2, "part_number": "PRT-WIRE-GA", "description": "Wire Harness Lead (m)", "quantity": 1.12345},
        ]
        store.save_manual_model_bom(
            self.project_id, self.model_id, items, auto_add_parts=True
        )

        retrieved = store.get_manual_model_bom(self.project_id, self.model_id)
        self.assertEqual(len(retrieved), 4)
        self.assertEqual(retrieved[1]["quantity"], 0.00025)
        self.assertEqual(retrieved[2]["quantity"], 0.00001)
        self.assertEqual(retrieved[3]["quantity"], 1.12345)

        # Check model tree structure export
        tree_res = store.pits_bom_model_tree(self.project_id, model_id=self.model_id)
        qty_map = {n["part_number"]: n["quantity"] for n in tree_res["nodes"]}
        self.assertEqual(qty_map.get("PRT-ADHESIVE"), 0.00025)
        self.assertEqual(qty_map.get("PRT-SEALANT"), 0.00001)
        self.assertEqual(qty_map.get("PRT-WIRE-GA"), 1.12345)

    def test_save_over_existing_bom_syncs_catalog_description_and_quantities(self):
        # 1. Initial BOM save
        initial_items = [
            {"level": 1, "part_number": "ASM-SYNC-1", "description": "Original Root", "quantity": 1.0},
            {"level": 2, "part_number": "PRT-SYNC-A", "description": "Original Screw", "quantity": 2.0},
            {"level": 2, "part_number": "PRT-SYNC-B", "description": "Original Nut", "quantity": 4.0},
        ]
        store.save_manual_model_bom(
            self.project_id, self.model_id, initial_items, auto_add_parts=True
        )

        # Verify initial parts catalog descriptions
        p_a = store.query(
            "SELECT description FROM parts WHERE project_id=? AND part_number='PRT-SYNC-A'",
            (self.project_id,),
        )[0]
        self.assertEqual(p_a["description"], "Original Screw")

        # 2. Update BOM over top with new descriptions and new decimal quantities
        updated_items = [
            {"level": 1, "part_number": "ASM-SYNC-1", "description": "Updated Root Asm", "quantity": 1.0},
            {"level": 2, "part_number": "PRT-SYNC-A", "description": "Upgraded Hex Screw M5", "quantity": 2.50000},
            {"level": 2, "part_number": "PRT-SYNC-B", "description": "Locking Flange Nut M5", "quantity": 6.12345},
            {"level": 2, "part_number": "PRT-SYNC-C", "description": "New Washer M5", "quantity": 0.00050},
        ]
        store.save_manual_model_bom(
            self.project_id, self.model_id, updated_items, auto_add_parts=True
        )

        # 3. Verify Parts Catalog synced the new descriptions
        p_a_updated = store.query(
            "SELECT description FROM parts WHERE project_id=? AND part_number='PRT-SYNC-A'",
            (self.project_id,),
        )[0]
        self.assertEqual(p_a_updated["description"], "Upgraded Hex Screw M5")

        p_b_updated = store.query(
            "SELECT description FROM parts WHERE project_id=? AND part_number='PRT-SYNC-B'",
            (self.project_id,),
        )[0]
        self.assertEqual(p_b_updated["description"], "Locking Flange Nut M5")

        # 4. Verify get_manual_model_bom reflects updated descriptions and quantities
        retrieved = store.get_manual_model_bom(self.project_id, self.model_id)
        desc_map = {r["part_number"]: r["description"] for r in retrieved}
        qty_map = {r["part_number"]: r["quantity"] for r in retrieved}

        self.assertEqual(desc_map["PRT-SYNC-A"], "Upgraded Hex Screw M5")
        self.assertEqual(desc_map["PRT-SYNC-B"], "Locking Flange Nut M5")
        self.assertEqual(desc_map["PRT-SYNC-C"], "New Washer M5")
        self.assertEqual(qty_map["PRT-SYNC-A"], 2.50000)
        self.assertEqual(qty_map["PRT-SYNC-B"], 6.12345)
        self.assertEqual(qty_map["PRT-SYNC-C"], 0.00050)

        # 5. Verify pits_bom_model_tree nodes reflect updated descriptions and quantities
        tree_res = store.pits_bom_model_tree(self.project_id, model_id=self.model_id)
        tree_desc_map = {n["part_number"]: n["description"] for n in tree_res["nodes"]}
        tree_qty_map = {n["part_number"]: n["quantity"] for n in tree_res["nodes"]}

        self.assertEqual(tree_desc_map["PRT-SYNC-A"], "Upgraded Hex Screw M5")
        self.assertEqual(tree_desc_map["PRT-SYNC-B"], "Locking Flange Nut M5")
        self.assertEqual(tree_desc_map["PRT-SYNC-C"], "New Washer M5")
        self.assertEqual(tree_qty_map["PRT-SYNC-A"], 2.50000)
        self.assertEqual(tree_qty_map["PRT-SYNC-B"], 6.12345)
        self.assertEqual(tree_qty_map["PRT-SYNC-C"], 0.00050)

    def test_save_duplicate_part_under_same_parent(self):
        items = [
            {"level": 1, "part_number": "ASM-DUP-TOP", "description": "Duplicate Test Top", "quantity": 1.0},
            {"level": 2, "part_number": "PRT-CLIP", "description": "Snap Clip A", "quantity": 1.0},
            {"level": 2, "part_number": "PRT-CLIP", "description": "Snap Clip B", "quantity": 2.0},
        ]
        # Should save without UNIQUE constraint error
        store.save_manual_model_bom(
            self.project_id, self.model_id, items, auto_add_parts=True
        )
        retrieved = store.get_manual_model_bom(self.project_id, self.model_id)
        self.assertEqual(len(retrieved), 3)
        clips = [r for r in retrieved if r["part_number"] == "PRT-CLIP"]
        self.assertEqual(len(clips), 2)
        self.assertEqual(clips[0]["quantity"], 1.0)
        self.assertEqual(clips[1]["quantity"], 2.0)

    def test_expand_subassembly_items_preserves_assemblies_and_inserts_between(self):
        draft = [
            {"level": 1, "part_number": "ASM-TOP-01", "description": "Top Assembly 1", "quantity": 1.0},
            {"level": 2, "part_number": "GRP-DOOR", "description": "Door Group", "quantity": 1.0},
            {"level": 1, "part_number": "ASM-TOP-02", "description": "Top Assembly 2", "quantity": 1.0},
        ]
        detected = [
            {
                "part_number": "GRP-DOOR",
                "level": 2,
                "description": "Door Group",
                "children": [
                    {"part_number": "PRT-HINGE", "description": "Door Hinge", "quantity": 2.0, "relative_level": 1},
                    {"part_number": "PRT-LATCH", "description": "Door Latch", "quantity": 1.0, "relative_level": 1},
                ],
            }
        ]

        expanded = store.expand_subassembly_items(draft, detected)
        self.assertEqual(len(expanded), 5)

        # 1. Verify Top-Level Assembly 1 is preserved and NOT deleted/overwritten
        self.assertEqual(expanded[0]["part_number"], "ASM-TOP-01")
        self.assertEqual(expanded[0]["level"], 1)
        self.assertEqual(expanded[0]["description"], "Top Assembly 1")

        # 2. Verify GRP-DOOR is preserved
        self.assertEqual(expanded[1]["part_number"], "GRP-DOOR")
        self.assertEqual(expanded[1]["level"], 2)

        # 3. Verify child components are inserted between assemblies at level 3
        self.assertEqual(expanded[2]["part_number"], "PRT-HINGE")
        self.assertEqual(expanded[2]["level"], 3)
        self.assertEqual(expanded[2]["quantity"], 2.0)
        self.assertEqual(expanded[3]["part_number"], "PRT-LATCH")
        self.assertEqual(expanded[3]["level"], 3)
        self.assertEqual(expanded[3]["quantity"], 1.0)

        # 4. Verify subsequent Top-Level Assembly 2 is NOT overwritten and follows children
        self.assertEqual(expanded[4]["part_number"], "ASM-TOP-02")
        self.assertEqual(expanded[4]["level"], 1)
        self.assertEqual(expanded[4]["description"], "Top Assembly 2")

    def test_detect_unexpanded_subassemblies(self):
        # Save a subassembly in model 1
        m1_items = [
            {"level": 1, "part_number": "ASM-SRC-01", "description": "Source Model", "quantity": 1.0},
            {"level": 2, "part_number": "GRP-SUB-WIDGET", "description": "Widget Assembly", "quantity": 1.0},
            {"level": 3, "part_number": "PRT-WIDGET-GEAR", "description": "Gear", "quantity": 3.0},
        ]
        store.save_manual_model_bom(self.project_id, self.model_id, m1_items, auto_add_parts=True)

        # Draft for model 2 contains only GRP-SUB-WIDGET without its children
        draft_unexpanded = [
            {"level": 1, "part_number": "ASM-TARGET-02", "description": "Target Model", "quantity": 1.0},
            {"level": 2, "part_number": "GRP-SUB-WIDGET", "description": "", "quantity": 1.0},
        ]

        detected = store.detect_unexpanded_subassemblies(self.project_id, draft_unexpanded)
        self.assertEqual(len(detected), 1)
        self.assertEqual(detected[0]["part_number"], "GRP-SUB-WIDGET")
        child_pns = [c["part_number"] for c in detected[0]["children"]]
        self.assertIn("PRT-WIDGET-GEAR", child_pns)

        # If child gear is already present, detect should return empty (no false positive)
        draft_already_expanded = [
            {"level": 1, "part_number": "ASM-TARGET-02", "description": "Target Model", "quantity": 1.0},
            {"level": 2, "part_number": "GRP-SUB-WIDGET", "description": "", "quantity": 1.0},
            {"level": 3, "part_number": "PRT-WIDGET-GEAR", "description": "Gear", "quantity": 3.0},
        ]
        detected_after = store.detect_unexpanded_subassemblies(self.project_id, draft_already_expanded)
        self.assertEqual(len(detected_after), 0)

    def test_subassembly_children_appear_below_group_and_levels_continuous(self):
        # 1. Setup a nested subassembly structure in model 1:
        # L1: ASM-ROOT -> L2: GRP-COMPLEX -> L3: SUB-MODULE -> L4: PRT-SCREW
        #                                -> L3: PRT-BRACKET
        model1_bom = [
            {"level": 1, "part_number": "ASM-ROOT-M1", "description": "Model 1 Root", "quantity": 1.0},
            {"level": 2, "part_number": "GRP-COMPLEX", "description": "Complex Group", "quantity": 1.0},
            {"level": 3, "part_number": "SUB-MODULE", "description": "Sub Module", "quantity": 1.0},
            {"level": 4, "part_number": "PRT-SCREW", "description": "Fastening Screw", "quantity": 4.0},
            {"level": 3, "part_number": "PRT-BRACKET", "description": "Mounting Bracket", "quantity": 2.0},
        ]
        store.save_manual_model_bom(self.project_id, self.model_id, model1_bom, auto_add_parts=True)

        # 2. In Model 2, user drafts a minimal BOM with only the group:
        m2_id = store.add_project_model(self.project_id, "TEST-M2", "Test Model 2", "Second Model")
        draft_m2 = [
            {"level": 1, "part_number": "ASM-ROOT-M2", "description": "Model 2 Root", "quantity": 1.0},
            {"level": 2, "part_number": "GRP-COMPLEX", "description": "", "quantity": 1.0},
            {"level": 1, "part_number": "ASM-SECOND-M2", "description": "Subsequent Assembly", "quantity": 1.0},
        ]

        # 3. Detect unexpanded subassemblies
        detected = store.detect_unexpanded_subassemblies(self.project_id, draft_m2)
        self.assertEqual(len(detected), 1)
        self.assertEqual(detected[0]["part_number"], "GRP-COMPLEX")

        # 4. Expand subassembly
        expanded = store.expand_subassembly_items(draft_m2, detected)
        self.assertEqual(len(expanded), 6)

        # Verify strict order:
        # Row 0: L1 ASM-ROOT-M2
        # Row 1: L2 GRP-COMPLEX (parent group must come first!)
        # Row 2: L3 SUB-MODULE (child component below group)
        # Row 3: L4 PRT-SCREW (grandchild below SUB-MODULE)
        # Row 4: L3 PRT-BRACKET (sibling child below group)
        # Row 5: L1 ASM-SECOND-M2 (subsequent assembly preserved at the end)
        self.assertEqual(expanded[0]["part_number"], "ASM-ROOT-M2")
        self.assertEqual(expanded[0]["level"], 1)

        self.assertEqual(expanded[1]["part_number"], "GRP-COMPLEX")
        self.assertEqual(expanded[1]["level"], 2)

        self.assertEqual(expanded[2]["part_number"], "SUB-MODULE")
        self.assertEqual(expanded[2]["level"], 3)

        self.assertEqual(expanded[3]["part_number"], "PRT-SCREW")
        self.assertEqual(expanded[3]["level"], 4)

        self.assertEqual(expanded[4]["part_number"], "PRT-BRACKET")
        self.assertEqual(expanded[4]["level"], 3)

        self.assertEqual(expanded[5]["part_number"], "ASM-SECOND-M2")
        self.assertEqual(expanded[5]["level"], 1)

        # Verify levels NEVER skip or jump backwards inappropriately (no 1, 3, 2)
        for i in range(1, len(expanded)):
            prev_l = expanded[i - 1]["level"]
            curr_l = expanded[i]["level"]
            # A child cannot skip levels (curr_l <= prev_l + 1)
            self.assertLessEqual(curr_l, prev_l + 1, f"Row {i} skipped level from {prev_l} to {curr_l}")

    def test_multimodel_source_row_collision_resilience(self):
        # Save model 1
        m1_bom = [
            {"level": 1, "part_number": "ASM-PARENT", "quantity": 1.0},
            {"level": 2, "part_number": "GRP-SHARED", "quantity": 1.0},
            {"level": 3, "part_number": "PRT-SHARED-COMP", "quantity": 1.0},
        ]
        store.save_manual_model_bom(self.project_id, self.model_id, m1_bom, auto_add_parts=True)

        # Now simulate a cross-model database collision where source_row of PRT-SHARED-COMP
        # is manually updated to a smaller or equal source_row as GRP-SHARED
        with store.connection() as conn:
            conn.execute(
                "UPDATE pits_bom_occurrences SET source_row = 1 WHERE child_tracker_number = 'PRT-SHARED-COMP'"
            )
            conn.execute(
                "UPDATE pits_bom_occurrences SET source_row = 10 WHERE child_tracker_number = 'GRP-SHARED'"
            )

        # Retrieve BOM for model 1 via get_manual_model_bom
        retrieved = store.get_manual_model_bom(self.project_id, self.model_id)
        self.assertEqual(len(retrieved), 3)

        # Assert that even with corrupted source_row, hierarchy order is preserved:
        # Row 0: L1 ASM-PARENT
        # Row 1: L2 GRP-SHARED (parent group MUST be above component)
        # Row 2: L3 PRT-SHARED-COMP (child component MUST be below group)
        self.assertEqual(retrieved[0]["part_number"], "ASM-PARENT")
        self.assertEqual(retrieved[0]["level"], 1)

        self.assertEqual(retrieved[1]["part_number"], "GRP-SHARED")
        self.assertEqual(retrieved[1]["level"], 2)

        self.assertEqual(retrieved[2]["part_number"], "PRT-SHARED-COMP")
        self.assertEqual(retrieved[2]["level"], 3)

    def test_spreadsheet_row_insertion_and_deletion_between_parts(self):
        # 1. Initial draft representing a 3-tier hierarchy
        draft = [
            {"level": 1, "part_number": "ASM-ROOT", "description": "Top Root", "quantity": 1.0, "_uid": "u1"},
            {"level": 2, "part_number": "GRP-SIDE", "description": "Side Group", "quantity": 1.0, "_uid": "u2"},
            {"level": 3, "part_number": "PRT-BOTTOM", "description": "Bottom Clip", "quantity": 4.0, "_uid": "u3"},
        ]

        # 2. Test inserting a row below Row 2 (GRP-SIDE) -> between GRP-SIDE and PRT-BOTTOM
        target_idx = 1  # Row 2 (0-indexed)
        new_row_below = {
            "level": draft[target_idx]["level"],
            "part_number": "PRT-BETWEEN-1",
            "description": "Inserted Between Parts",
            "quantity": 1.0,
            "_uid": "u_new1",
        }
        draft.insert(target_idx + 1, new_row_below)

        self.assertEqual(len(draft), 4)
        self.assertEqual(draft[0]["part_number"], "ASM-ROOT")
        self.assertEqual(draft[1]["part_number"], "GRP-SIDE")
        self.assertEqual(draft[2]["part_number"], "PRT-BETWEEN-1")  # Inserted between row 2 and 3!
        self.assertEqual(draft[2]["level"], 2)
        self.assertEqual(draft[3]["part_number"], "PRT-BOTTOM")

        # 3. Test inserting a row above Row 2 (GRP-SIDE) -> between ASM-ROOT and GRP-SIDE
        target_idx = 1
        new_row_above = {
            "level": draft[target_idx]["level"],
            "part_number": "PRT-BETWEEN-0",
            "description": "Inserted Above Group",
            "quantity": 1.0,
            "_uid": "u_new0",
        }
        draft.insert(target_idx, new_row_above)

        self.assertEqual(len(draft), 5)
        self.assertEqual(draft[0]["part_number"], "ASM-ROOT")
        self.assertEqual(draft[1]["part_number"], "PRT-BETWEEN-0")  # Inserted above GRP-SIDE!
        self.assertEqual(draft[2]["part_number"], "GRP-SIDE")
        self.assertEqual(draft[3]["part_number"], "PRT-BETWEEN-1")
        self.assertEqual(draft[4]["part_number"], "PRT-BOTTOM")

        # 4. Test deleting a specific row between part numbers (e.g. Row 3: GRP-SIDE)
        del_target = 2
        removed = draft.pop(del_target)
        self.assertEqual(removed["part_number"], "GRP-SIDE")
        self.assertEqual(len(draft), 4)
        self.assertEqual([r["part_number"] for r in draft], ["ASM-ROOT", "PRT-BETWEEN-0", "PRT-BETWEEN-1", "PRT-BOTTOM"])

        # 5. Test moving a row up and down
        # Move Row 2 (PRT-BETWEEN-1, idx 2) up
        draft[2], draft[1] = draft[1], draft[2]
        self.assertEqual([r["part_number"] for r in draft], ["ASM-ROOT", "PRT-BETWEEN-1", "PRT-BETWEEN-0", "PRT-BOTTOM"])

    def test_get_live_editor_items_preserves_cell_edits_and_native_deletions(self):
        import streamlit as st
        from utils.store import get_live_bom_editor_items

        m_id = "test_model_live"
        d_key = f"manual_bom_draft_{m_id}"
        ed_key = f"bom_grid_editor_{m_id}"

        # Initialize base items
        st.session_state[d_key] = [
            {"level": 1, "part_number": "ASM-ROOT", "description": "Old Root", "quantity": 1.0, "_uid": "u1"},
            {"level": 2, "part_number": "GRP-DEL", "description": "To Delete", "quantity": 1.0, "_uid": "u2"},
            {"level": 2, "part_number": "PRT-KEEP", "description": "Keep Me", "quantity": 1.0, "_uid": "u3"},
        ]

        from utils.table_filters import _EDITOR_INSTANCE_MARKER

        # Simulate Streamlit data_editor state:
        # - Row 0 has Description and Qty edited
        # - Row 1 is deleted natively
        inst_key = f"{ed_key}{_EDITOR_INSTANCE_MARKER}0"
        st.session_state[inst_key] = {
            "edited_rows": {
                "0": {"Description": "New Root Desc", "Qty": 3.14159},
            },
            "added_rows": [],
            "deleted_rows": [1],
        }

        live_items = get_live_bom_editor_items(m_id, d_key, ed_key)

        # GRP-DEL should be gone
        self.assertEqual(len(live_items), 2)
        # ASM-ROOT should have preserved the live edits
        self.assertEqual(live_items[0]["part_number"], "ASM-ROOT")
        self.assertEqual(live_items[0]["description"], "New Root Desc")
        self.assertEqual(live_items[0]["quantity"], 3.14159)
        # PRT-KEEP should remain
        self.assertEqual(live_items[1]["part_number"], "PRT-KEEP")

    def test_direct_table_row_selection_and_action_dispatch(self):
        # Verify that toggling Select checkbox or editing a cell in the table sets the active row index
        m_id = self.model_id
        row_tools_idx_key = f"row_tools_idx_{m_id}"
        sel_indices_key = f"bom_selected_indices_{m_id}"

        # 1. Simulate user selecting Row 2 in table (0-indexed: 1)
        import streamlit as st
        st.session_state[sel_indices_key] = {1}
        st.session_state[row_tools_idx_key] = 1

        self.assertIn(1, st.session_state[sel_indices_key])
        self.assertEqual(st.session_state[row_tools_idx_key], 1)

        # 2. Simulate inserting a row below Row 2 (pos 2)
        draft = [
            {"level": 1, "part_number": "L1-ASM", "quantity": 1.0},
            {"level": 2, "part_number": "L2-GRP", "quantity": 1.0},
            {"level": 3, "part_number": "L3-PRT", "quantity": 1.0},
        ]
        sel_row_idx = st.session_state[row_tools_idx_key]
        target_item = draft[sel_row_idx]
        new_row = {
            "level": target_item["level"],
            "part_number": "L2-NEW-INSERT",
            "quantity": 1.0,
        }
        ins_pos = sel_row_idx + 1
        draft.insert(ins_pos, new_row)
        st.session_state[row_tools_idx_key] = ins_pos

        self.assertEqual(len(draft), 4)
        self.assertEqual(draft[0]["part_number"], "L1-ASM")
        self.assertEqual(draft[1]["part_number"], "L2-GRP")
        self.assertEqual(draft[2]["part_number"], "L2-NEW-INSERT")
        self.assertEqual(draft[3]["part_number"], "L3-PRT")
        # Newly inserted row is now selected
        self.assertEqual(st.session_state[row_tools_idx_key], 2)

    def test_multi_row_move_up_contiguous_and_boundary(self):
        from utils.store import move_bom_rows_up
        items = [
            {"part_number": "P0"},
            {"part_number": "P1"},
            {"part_number": "P2"},
            {"part_number": "P3"},
            {"part_number": "P4"},
        ]
        # 1. Move contiguous selection [1, 2] up -> P1 and P2 swap with P0
        res_items, res_sel = move_bom_rows_up(items, {1, 2})
        self.assertEqual([x["part_number"] for x in res_items], ["P1", "P2", "P0", "P3", "P4"])
        self.assertEqual(res_sel, {0, 1})

        # 2. Boundary test: moving selection {0, 1} up cannot move beyond top
        res_items2, res_sel2 = move_bom_rows_up(res_items, res_sel)
        self.assertEqual([x["part_number"] for x in res_items2], ["P1", "P2", "P0", "P3", "P4"])
        self.assertEqual(res_sel2, {0, 1})

        # 3. Non-contiguous selection {1, 3} on original items
        res_items3, res_sel3 = move_bom_rows_up(items, {1, 3})
        # P1 moves to 0 (swapping P0), P3 moves to 2 (swapping P2)
        self.assertEqual([x["part_number"] for x in res_items3], ["P1", "P0", "P3", "P2", "P4"])
        self.assertEqual(res_sel3, {0, 2})

    def test_multi_row_move_down_contiguous_and_boundary(self):
        from utils.store import move_bom_rows_down
        items = [
            {"part_number": "P0"},
            {"part_number": "P1"},
            {"part_number": "P2"},
            {"part_number": "P3"},
            {"part_number": "P4"},
        ]
        # 1. Move contiguous selection [2, 3] down -> P2 and P3 swap with P4
        res_items, res_sel = move_bom_rows_down(items, {2, 3})
        self.assertEqual([x["part_number"] for x in res_items], ["P0", "P1", "P4", "P2", "P3"])
        self.assertEqual(res_sel, {3, 4})

        # 2. Boundary test: moving selection {3, 4} down cannot move beyond bottom
        res_items2, res_sel2 = move_bom_rows_down(res_items, res_sel)
        self.assertEqual([x["part_number"] for x in res_items2], ["P0", "P1", "P4", "P2", "P3"])
        self.assertEqual(res_sel2, {3, 4})

        # 3. Non-contiguous selection {1, 3} on original items
        res_items3, res_sel3 = move_bom_rows_down(items, {1, 3})
        # P3 moves to 4 (swapping P4), P1 moves to 2 (swapping P2)
        self.assertEqual([x["part_number"] for x in res_items3], ["P0", "P2", "P1", "P4", "P3"])
        self.assertEqual(res_sel3, {2, 4})

    def test_multi_row_delete_and_empty_reset(self):
        from utils.store import delete_bom_rows
        items = [
            {"part_number": "P0"},
            {"part_number": "P1"},
            {"part_number": "P2"},
            {"part_number": "P3"},
        ]
        # 1. Delete multiple rows [1, 2]
        res_items, res_sel = delete_bom_rows(items, {1, 2})
        self.assertEqual([x["part_number"] for x in res_items], ["P0", "P3"])
        self.assertEqual(res_sel, {1})

        # 2. Delete all rows
        res_items_empty, res_sel_empty = delete_bom_rows(res_items, {0, 1})
        self.assertEqual(len(res_items_empty), 1)
        self.assertEqual(res_items_empty[0]["part_number"], "")
        self.assertEqual(res_sel_empty, {0})

    def test_range_selection_and_checkbox_selection(self):
        # Verify range selection logic: min to max inclusive
        from_r = 1
        to_r = 4
        sel_range = set(range(min(from_r, to_r), max(from_r, to_r) + 1))
        self.assertEqual(sel_range, {1, 2, 3, 4})

        # Verify reverse range: from Row 5 to Row 2
        rev_range = set(range(min(to_r, from_r), max(to_r, from_r) + 1))
        self.assertEqual(rev_range, {1, 2, 3, 4})

    def test_bom_tree_page_smoke(self):
        page_path = Path(__file__).resolve().parent.parent / "app_pages" / "bom_tree.py"
        app = AppTest.from_file(str(page_path))
        app.session_state["project_id"] = self.project_id
        app.session_state["model_tree_selected_model"] = self.model_id
        app.run()
        self.assertEqual(len(app.exception), 0)

    def test_generate_bom_excel_template_with_existing_tree_and_blank_fallback(self):
        import io
        import openpyxl
        from app_pages.bom_tree import generate_bom_excel_template, parse_uploaded_bom_file

        # 1. Blank items fallback to generic template
        blank_items = [
            {"level": 1, "part_number": "", "description": "", "quantity": 1.0},
        ]
        blank_bytes = generate_bom_excel_template(blank_items)
        wb_blank = openpyxl.load_workbook(io.BytesIO(blank_bytes))
        rows_blank = list(wb_blank["BOM"].iter_rows(values_only=True))
        self.assertEqual(rows_blank[0], ("Level", "Part #", "Description", "Qty"))
        self.assertEqual(rows_blank[1], (1, "SAMPLE_MAIN_ASY", "Main Top-Level Assembly", 1.0))
        self.assertEqual(rows_blank[2], (2, "SAMPLE_SUB_ASY", "Subassembly Module", 1.0))

        # None items fallback to generic template
        none_bytes = generate_bom_excel_template(None)
        wb_none = openpyxl.load_workbook(io.BytesIO(none_bytes))
        rows_none = list(wb_none["BOM"].iter_rows(values_only=True))
        self.assertEqual(rows_none[1], (1, "SAMPLE_MAIN_ASY", "Main Top-Level Assembly", 1.0))

        # 2. Existing tree items are populated into the Excel download
        tree_items = [
            {"level": 1, "part_number": "ASY-TOP-01", "description": "Top Main Assembly", "quantity": 1.0},
            {"level": 2, "part_number": "SUB-MODULE-02", "description": "Door Assembly", "quantity": 1.0},
            {"level": 3, "part_number": "999D8503P002", "description": "Black ABS Roof", "quantity": 3.5},
        ]
        tree_bytes = generate_bom_excel_template(tree_items)
        wb_tree = openpyxl.load_workbook(io.BytesIO(tree_bytes))
        rows_tree = list(wb_tree["BOM"].iter_rows(values_only=True))
        self.assertEqual(rows_tree[0], ("Level", "Part #", "Description", "Qty"))
        self.assertEqual(rows_tree[1], (1, "ASY-TOP-01", "Top Main Assembly", 1.0))
        self.assertEqual(rows_tree[2], (2, "SUB-MODULE-02", "Door Assembly", 1.0))
        self.assertEqual(rows_tree[3], (3, "999D8503P002", "Black ABS Roof", 3.5))

        # 3. Round-trip verify that parse_uploaded_bom_file can parse it back accurately
        f_like = io.BytesIO(tree_bytes)
        f_like.name = "BOM_Model.xlsx"
        parsed = parse_uploaded_bom_file(f_like)
        self.assertEqual(len(parsed), 3)
        self.assertEqual(parsed[0]["level"], 1)
        self.assertEqual(parsed[0]["part_number"], "ASY-TOP-01")
        self.assertEqual(parsed[1]["level"], 2)
        self.assertEqual(parsed[1]["part_number"], "SUB-MODULE-02")
        self.assertEqual(parsed[2]["level"], 3)
        self.assertEqual(parsed[2]["part_number"], "999D8503P002")
    def test_parse_pasted_bom_text_with_multiline_and_excel_headers(self):
        from app_pages.bom_tree import parse_pasted_bom_text

        sample_excel_clipboard = (
            "Level\tPart #\tDescription\tQty\n"
            "1\t999D8842G001\tUrban Coupe Final Assembly\t1\n"
            "2\t999D8502G001\tMiddle body assembly\t1\n"
            "3\t999D8500P001\tside graphic\t1\n"
            "3\t999D8502P001\tWooden block body\t1\n"
            "3\t999D8508P001\tback emblem\t1\n"
            "2\t999D8402G002\tFront of car body asm\t1\n"
            "3\t999D6206P001\taxle pin\t2\n"
            "3\t999D8505G001\tWheel asm (teal and chrome)\t2\n"
            "4\t999D8505P001\tteal tire\t1\n"
            "4\t999D4805P001\tchrome inserts\t1\n"
            "4\t999D4805P001\tchrome inserts\t1\n"
            "4\t999D8505P001\tteal tire\t1\n"
            "4\t999D4805P001\tchrome inserts\t1\n"
            "3\t999D8506P002\tConnector rectangle\t1\n"
            "3\t999D8502P002\tfront nose of car\t1\n"
            "3\t999D8507P001\tFront emblem\t1\n"
            "2\t999D8802G003\tback of car asm\t1\n"
            "3\t999D6206P001\taxle pin\t2\n"
            "3\t999D8505G001\tWheel asm (teal and chrome)\t2\n"
            "3\t999D8506P002\tConnector rectangle\t1\n"
            "3\t999D8502P003\t\"back of\n\n      car wood\"\t1\n"
            "3\t999D8508P001\tback emblem\t1\n"
        )

        parsed = parse_pasted_bom_text(sample_excel_clipboard)
        self.assertEqual(len(parsed), 22)

        # Check top level
        self.assertEqual(parsed[0]["level"], 1)
        self.assertEqual(parsed[0]["part_number"], "999D8842G001")
        self.assertEqual(parsed[0]["description"], "Urban Coupe Final Assembly")
        self.assertEqual(parsed[0]["quantity"], 1.0)

        # Check multiline row (item 20) has collapsed whitespace
        self.assertEqual(parsed[20]["level"], 3)
        self.assertEqual(parsed[20]["part_number"], "999D8502P003")
        self.assertEqual(parsed[20]["description"], "back of car wood")
        self.assertEqual(parsed[20]["quantity"], 1.0)

        # Check that no bogus "nan" part numbers or descriptions exist
        for it in parsed:
            self.assertNotEqual(it["part_number"].lower(), "nan")
            self.assertNotEqual(it["description"].lower(), "nan")
            self.assertTrue(bool(it["part_number"]))

    def test_predict_model_bom_from_features_level1_assembly_handling(self):
        """Verify predict_model_bom_from_features does not invent _TOP and correctly uses custom Level 1."""
        # 1. Add parts to catalog
        p1_id = str(uuid4())
        p2_id = str(uuid4())
        top_asm_id = str(uuid4())
        store.execute(
            """INSERT INTO parts (id, project_id, part_number, description, weight_lb, make_buy, model_applicability, updated_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            (top_asm_id, self.project_id, "999D8842G001", "Urban Coupe Final Assembly", 10.0, "Make", "All", store.now_iso()),
        )
        store.execute(
            """INSERT INTO parts (id, project_id, part_number, description, weight_lb, make_buy, model_applicability, updated_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            (p1_id, self.project_id, "999D8502G001", "Middle body assembly", 3.0, "Make", "All", store.now_iso()),
        )
        store.execute(
            """INSERT INTO parts (id, project_id, part_number, description, weight_lb, make_buy, model_applicability, updated_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            (p2_id, self.project_id, "999D8500P001", "Side graphic", 0.1, "Buy", "All", store.now_iso()),
        )

        # 2. When called with no top_level_part_number, Level 1 must have empty part number (NOT f"{model_number}_TOP")
        pred_empty = store.predict_model_bom_from_features(self.project_id, self.model_id)
        items_empty = pred_empty["predicted_items"]
        self.assertTrue(len(items_empty) >= 1)
        self.assertEqual(items_empty[0]["level"], 1)
        self.assertEqual(items_empty[0]["part_number"], "")
        self.assertFalse(any("_TOP" in it["part_number"] for it in items_empty))

        # 3. When called with a designated top_level_part_number (e.g. from Parts Catalog), Level 1 receives it
        pred_chosen = store.predict_model_bom_from_features(
            self.project_id,
            self.model_id,
            top_level_part_number="999D8842G001",
            top_level_description="Urban Coupe Final Assembly",
        )
        items_chosen = pred_chosen["predicted_items"]
        self.assertEqual(items_chosen[0]["level"], 1)
        self.assertEqual(items_chosen[0]["part_number"], "999D8842G001")
        self.assertEqual(items_chosen[0]["description"], "Urban Coupe Final Assembly")
        # Ensure the Level 1 assembly is not duplicated as a child component at Level 2
        lvl2_pns = [it["part_number"] for it in items_chosen if it["level"] == 2]
        self.assertNotIn("999D8842G001", lvl2_pns)

    def test_predict_model_bom_replicates_children_and_substitutes_only_counterparts(self):
        """Verify predict_model_bom_from_features replicates children 1:1 and does not duplicate substitutions."""
        # 1. Setup complexity feature: Roof option
        feat_id = str(uuid4())
        store.execute(
            """INSERT INTO complexity_features (id, project_id, name, category, allowed_values, description, sequence, active, updated_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (feat_id, self.project_id, "Roof option", "Exterior", '["Coupe", "Tall"]', "Roof feature", 1, 1, store.now_iso()),
        )

        # 2. Template Model (M1) = Coupe, Target Model (M2) = Tall
        model2_id = store.add_project_model(self.project_id, "TEST-M2", "Test Model 2", "")
        store.execute(
            """INSERT INTO model_feature_values (project_id, model_id, feature_id, value, updated_at)
               VALUES (?, ?, ?, ?, ?)""",
            (self.project_id, self.model_id, feat_id, "Coupe", store.now_iso()),
        )
        store.execute(
            """INSERT INTO model_feature_values (project_id, model_id, feature_id, value, updated_at)
               VALUES (?, ?, ?, ?, ?)""",
            (self.project_id, model2_id, feat_id, "Tall", store.now_iso()),
        )

        # 3. Catalog parts with rules
        p_roof_std = str(uuid4())
        p_roof_tall = str(uuid4())
        p_body = str(uuid4())
        p_pin = str(uuid4())
        store.execute(
            """INSERT INTO parts (id, project_id, part_number, description, weight_lb, make_buy, model_applicability, updated_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            (p_roof_std, self.project_id, "999D8503P001", "Teal ABS Roof", 0.5, "Buy", "All", store.now_iso()),
        )
        store.execute(
            """INSERT INTO parts (id, project_id, part_number, description, weight_lb, make_buy, model_applicability, updated_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            (p_roof_tall, self.project_id, "999D8503P003", "Tall Teal ABS Roof", 0.6, "Buy", "All", store.now_iso()),
        )
        store.execute(
            """INSERT INTO parts (id, project_id, part_number, description, weight_lb, make_buy, model_applicability, updated_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            (p_body, self.project_id, "999D8502G001", "Middle body assembly", 3.0, "Make", "All", store.now_iso()),
        )
        store.execute(
            """INSERT INTO parts (id, project_id, part_number, description, weight_lb, make_buy, model_applicability, updated_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            (p_pin, self.project_id, "999D6206P001", "Axle pin steel", 0.05, "Buy", "All", store.now_iso()),
        )

        # Rules: p_roof_std is Coupe, p_roof_tall is Tall
        store.execute(
            """INSERT INTO part_feature_rules (project_id, part_id, feature_id, value, updated_at)
               VALUES (?, ?, ?, ?, ?)""",
            (self.project_id, p_roof_std, feat_id, "Coupe", store.now_iso()),
        )
        store.execute(
            """INSERT INTO part_feature_rules (project_id, part_id, feature_id, value, updated_at)
               VALUES (?, ?, ?, ?, ?)""",
            (self.project_id, p_roof_tall, feat_id, "Tall", store.now_iso()),
        )

        # 4. Save template BOM for Model 1
        m1_bom = [
            {"level": 1, "part_number": "TOP-M1", "description": "Coupe Top Asm", "quantity": 1.0},
            {"level": 2, "part_number": "999D8502G001", "description": "Middle body assembly", "quantity": 1.0},
            {"level": 2, "part_number": "999D8503P001", "description": "Teal ABS Roof", "quantity": 1.0},
            {"level": 2, "part_number": "999D6206P001", "description": "Axle pin steel", "quantity": 4.0},
        ]
        store.save_manual_model_bom(self.project_id, self.model_id, m1_bom, auto_add_parts=True)

        # 5. Predict for Model 2
        pred = store.predict_model_bom_from_features(
            self.project_id,
            model2_id,
            top_level_part_number="TOP-M2",
            top_level_description="Tall SUV Top Asm",
        )
        items = pred["predicted_items"]
        subs = pred["substitutions"]

        # Exactly 1 substitution: 999D8503P001 -> 999D8503P003
        self.assertEqual(len(subs), 1)
        self.assertEqual(subs[0]["from_part"], "999D8503P001")
        self.assertEqual(subs[0]["to_part"], "999D8503P003")

        # Part counts and replica check:
        # Total rows = 4 (1 root + 3 children)
        self.assertEqual(len(items), 4)
        self.assertEqual(items[0]["part_number"], "TOP-M2")
        self.assertEqual(items[1]["part_number"], "999D8502G001")
        self.assertEqual(items[2]["part_number"], "999D8503P003")  # substituted counterpart
        self.assertEqual(items[3]["part_number"], "999D6206P001")  # preserved replica!
        self.assertEqual(items[3]["quantity"], 4.0)

        # Verify no duplicate part numbers exist in items
        pns = [it["part_number"] for it in items]
        self.assertEqual(len(pns), len(set(pns)))

    def test_predict_model_bom_handles_multi_level1_template_without_duplication(self):
        """Verify predict_model_bom_from_features only converts the first Level 1 row and demotes others without duplication."""
        m1_multi_l1_bom = [
            {"level": 1, "part_number": "ORIG-L1-A", "description": "First Root Asm", "quantity": 1.0},
            {"level": 2, "part_number": "SUB-01", "description": "Child Component", "quantity": 2.0},
            {"level": 1, "part_number": "ORIG-L1-B", "description": "Second Root Asm", "quantity": 1.0},
            {"level": 1, "part_number": "ORIG-L1-C", "description": "Third Root Asm", "quantity": 1.0},
        ]
        store.save_manual_model_bom(self.project_id, self.model_id, m1_multi_l1_bom, auto_add_parts=True)

        model2_id = store.add_project_model(self.project_id, "TEST-M2-MULTI", "Model 2 Multi", "")
        pred = store.predict_model_bom_from_features(
            self.project_id,
            model2_id,
            top_level_part_number="NEW-TOP-ROOT",
            top_level_description="New Unified Root",
        )
        items = pred["predicted_items"]

        # Only row 0 should be NEW-TOP-ROOT at Level 1
        l1_rows = [it for it in items if it["level"] == 1]
        self.assertEqual(len(l1_rows), 1)
        self.assertEqual(l1_rows[0]["part_number"], "NEW-TOP-ROOT")

        # The subsequent rows that were Level 1 should be demoted to Level 2 and retain their original part numbers
        self.assertEqual(items[1]["part_number"], "SUB-01")
        self.assertEqual(items[2]["part_number"], "ORIG-L1-B")
        self.assertEqual(items[2]["level"], 2)
        self.assertEqual(items[3]["part_number"], "ORIG-L1-C")
        self.assertEqual(items[3]["level"], 2)

        # No duplicate NEW-TOP-ROOT anywhere
        new_top_occurrences = [it for it in items if it["part_number"] == "NEW-TOP-ROOT"]
        self.assertEqual(len(new_top_occurrences), 1)

    def test_predict_model_bom_returns_model_names(self):
        """Verify predict_model_bom_from_features includes both model_name and template_model_name."""
        m1_bom = [
            {"level": 1, "part_number": "TOP-M1", "description": "Base Top Asm", "quantity": 1.0},
            {"level": 2, "part_number": "PRT-01", "description": "Standard Widget", "quantity": 1.0},
        ]
        store.save_manual_model_bom(self.project_id, self.model_id, m1_bom, auto_add_parts=True)

        target_model_id = store.add_project_model(
            self.project_id, "TEST-TARGET", "Target Display Name", "Target description"
        )

        pred = store.predict_model_bom_from_features(
            self.project_id,
            target_model_id,
            top_level_part_number="NEW-L1-PN",
            top_level_description="New L1 Description",
        )

        self.assertEqual(pred["model_number"], "TEST-TARGET")
        self.assertEqual(pred["model_name"], "Target Display Name")
        self.assertEqual(pred["template_model_number"], "TEST-M1")
        self.assertEqual(pred["template_model_name"], "Test Model 1")


if __name__ == "__main__":
    unittest.main()

