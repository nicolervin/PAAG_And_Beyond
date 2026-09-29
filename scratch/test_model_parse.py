import openpyxl

wb = openpyxl.load_workbook('scratch/pits_sample.xlsm', data_only=True, read_only=True)
sheet = wb['BOM']

row1 = next(sheet.iter_rows(values_only=True))

models_to_test = [
    (20, row1[20]),
    (21, row1[21]),
    (22, row1[22]),
    (23, row1[23]),
    (24, row1[24]),
]

for col_idx, m_name in models_to_test:
    nodes = []
    stack = {}
    for r_idx, row in enumerate(sheet.iter_rows(values_only=True)):
        if r_idx < 8:
            continue
        u = row[col_idx]
        try:
            if u is None or float(u) <= 0:
                continue
            qty = float(u)
        except (ValueError, TypeError):
            continue
        
        depth = None
        for l in range(3, 14):
            if row[l] is not None and str(row[l]).strip():
                depth = l - 2
                break
        if not depth:
            continue
        
        node = {'depth': depth, 'part': row[1], 'desc': row[2], 'qty': qty, 'children': []}
        parent = stack.get(depth - 1)
        if parent:
            parent['children'].append(node)
        stack[depth] = node
        stack = {k: v for k, v in stack.items() if k <= depth}
        nodes.append(node)
        
    roots = [n for n in nodes if n['depth'] == 1]
    print(f"Model {m_name} (Col {col_idx}): Total {len(nodes)} parts, {len(roots)} roots")
    for r in roots:
        print(f"  Root: {r['part']} - {r['desc']}, L2 children: {len(r['children'])}")
